# -*- coding: utf-8 -*-
"""TRENDMECCA 2026 4분기 운영계획 — 보고용 Streamlit 앱

실행:  streamlit run app.py
데이터: data/ 폴더의 목표매출 엑셀(파일명에 '목표')과 상품등급 엑셀(파일명에 '등급')을 읽습니다.
       파일을 교체하면 됩니다.
"""
import io
import re
from datetime import date
from pathlib import Path

import pandas as pd
import plotly.graph_objects as go
import streamlit as st
import streamlit.components.v1 as components
from openpyxl import load_workbook

BASE = Path(__file__).parent
DATA = BASE / "data"

INK, INK2, SOFT, MUTED, RED, RED_L, MID, LINE = "#141414", "#262626", "#F3F3F3", "#6E6E6E", "#C8102E", "#FF5A6E", "#9A9A9A", "#DADADA"
FONT = "Pretendard, 'Malgun Gothic', '맑은 고딕', 'Apple SD Gothic Neo', sans-serif"

st.set_page_config(page_title="TRENDMECCA 2026 Q4 운영계획", page_icon="📊", layout="wide", initial_sidebar_state="expanded")

# ------------------------------------------------------------------ style
st.markdown(f"""
<link rel="stylesheet" href="https://cdn.jsdelivr.net/gh/orioncactus/pretendard@v1.3.9/dist/web/static/pretendard.min.css">
<style>
html, body, [class*="css"], .stMarkdown, .stDataFrame, button, input, textarea, select {{ font-family: {FONT} !important; }}
#MainMenu, footer {{ visibility: hidden; }}
[data-testid="stSidebarCollapsedControl"], [data-testid="stExpandSidebarButton"] {{ visibility: visible !important; display: flex !important; z-index: 1000; }}
[data-testid="stSidebarCollapsedControl"] button, [data-testid="stExpandSidebarButton"] {{ background: #141414; color: #fff; border-radius: 10px; }}
.block-container {{ padding-top: 0; padding-bottom: 0; max-width: 1320px; }}
[data-testid="stMain"] {{ scroll-behavior: smooth; }}
div[class*="st-key-sec"] {{ min-height: 100vh; padding-top: 2.6rem; padding-bottom: 1rem; box-sizing: border-box; }}
.anchor {{ position: relative; top: -2.6rem; height: 0; }}
.toc a {{ display: block; padding: 9px 12px; margin-bottom: 4px; border-radius: 10px; color: #E8E8E8 !important; text-decoration: none !important; font-size: 14px; }}
.toc a:hover {{ background: {INK2}; }}
.toc a:before {{ content: ""; display: inline-block; width: 7px; height: 7px; border-radius: 50%; background: {RED}; margin-right: 10px; vertical-align: middle; }}
.toc-tip {{ color: #6E6E6E; font-size: 11.5px; line-height: 1.6; margin-top: 18px; padding-left: 12px; }}
section[data-testid="stSidebar"] {{ background: {INK}; }}
section[data-testid="stSidebar"] * {{ color: #E8E8E8; }}
section[data-testid="stSidebar"] [data-testid="stFileUploaderDropzone"] {{ background: {INK2}; border: 1px dashed #444; }}
section[data-testid="stSidebar"] [data-testid="stFileUploaderDropzone"] * {{ color: #BDBDBD !important; }}
section[data-testid="stSidebar"] [role="radiogroup"] label {{ padding: 8px 10px; border-radius: 10px; margin-bottom: 2px; width: 100%; }}
section[data-testid="stSidebar"] [role="radiogroup"] label:hover {{ background: {INK2}; }}
.kicker {{ color: {RED}; font-weight: 800; font-size: 13px; letter-spacing: .14em; margin-bottom: 4px; }}
.h1 {{ font-size: 28px; font-weight: 800; color: {INK}; letter-spacing: -.02em; line-height: 1.2; margin: 0 0 6px; }}
.sub {{ font-size: 14px; color: {MUTED}; margin-bottom: 14px; }}
.sec {{ font-size: 16px; font-weight: 800; color: {INK}; margin: 16px 0 8px; }}
.card {{ background: {SOFT}; border-radius: 14px; padding: 16px 20px; height: 100%; }}
.card.dark {{ background: {INK}; color: #fff; }}
.pill {{ display: inline-block; padding: 4px 14px; border-radius: 999px; font-weight: 800; font-size: 13px; color: #fff; background: {INK}; }}
.pill.red {{ background: {RED}; }}
.pill.mid {{ background: {MID}; }}
.big {{ font-size: 48px; font-weight: 800; letter-spacing: -.03em; line-height: 1; margin: 12px 0 10px; }}
.big small {{ font-size: 26px; }}
.track {{ height: 10px; border-radius: 6px; background: #DCDCDC; overflow: hidden; }}
.dark .track {{ background: #3A3A3A; }}
.fill {{ height: 100%; border-radius: 6px; background: {RED}; }}
.stats {{ display: flex; gap: 18px; margin-top: 22px; }}
.stats div {{ flex: 1; }}
.stats .l {{ font-size: 12px; color: {MUTED}; }}
.dark .stats .l {{ color: #A8A8A8; }}
.stats .v {{ font-size: 24px; font-weight: 800; margin-top: 4px; }}
.stats .v span {{ font-size: 13px; font-weight: 500; margin-left: 2px; }}
.stats .v.r {{ color: {RED}; }}
.dark .stats .v.r {{ color: {RED_L}; }}
.hero {{ background: {INK}; border-radius: 20px; padding: 32px 44px; position: relative; overflow: hidden; color: #fff; }}
.hero .q4 {{ position: absolute; right: 30px; top: -40px; font-size: 220px; font-weight: 900; color: #1F1F1F; line-height: 1; letter-spacing: -.04em; }}
.hero .t {{ position: relative; font-size: 36px; font-weight: 800; line-height: 1.18; letter-spacing: -.02em; margin: 14px 0 18px; }}
.hero .k {{ position: relative; color: {RED}; font-weight: 800; letter-spacing: .3em; font-size: 13px; }}
.hero .s {{ position: relative; color: #B5B5B5; font-size: 15px; }}
.agenda {{ display: flex; align-items: center; gap: 22px; background: {SOFT}; border-radius: 12px; padding: 10px 22px; margin-bottom: 7px; }}
.agenda .n {{ color: {RED}; font-weight: 800; font-size: 22px; width: 36px; }}
.agenda .t {{ font-weight: 800; font-size: 17px; width: 190px; }}
.agenda .d {{ color: {MUTED}; font-size: 14px; }}
.dir {{ display: flex; gap: 18px; align-items: flex-start; }}
.dot {{ flex: 0 0 44px; height: 44px; border-radius: 50%; background: {RED}; color: #fff; display: flex; align-items: center; justify-content: center; font-size: 22px; }}
.dir .t {{ font-weight: 800; font-size: 16px; margin-bottom: 6px; }}
.dir .d {{ color: {MUTED}; font-size: 14px; line-height: 1.6; }}
.num {{ float: right; color: #C9C9C9; font-weight: 800; font-size: 26px; }}
.steps {{ display: flex; gap: 10px; align-items: center; margin: 0 0 12px; }}
.step {{ flex: 1; background: {SOFT}; border-radius: 999px; padding: 8px 16px; font-weight: 800; font-size: 14px; }}
.step b {{ color: {RED}; margin-right: 8px; }}
.step.on {{ background: {INK}; color: #fff; }}
.step.on b {{ color: {RED_L}; }}
.chev {{ color: {MID}; font-weight: 800; }}
.band {{ display: flex; align-items: center; gap: 18px; background: {SOFT}; border-radius: 14px; padding: 12px 20px; margin-bottom: 9px; }}
.band.dark {{ background: {INK}; color: #fff; }}
.band .b {{ flex: 0 0 84px; text-align: center; padding: 10px 0; border-radius: 999px; color: #fff; font-weight: 800; font-size: 16px; }}
.band .m {{ flex: 0 0 190px; }}
.band .m .n {{ font-weight: 800; font-size: 17px; }}
.band .m .s {{ font-size: 12.5px; color: {MUTED}; line-height: 1.5; margin-top: 4px; }}
.band.dark .m .s {{ color: #BDBDBD; }}
.band .a {{ flex: 1; }}
.band .a .x {{ font-weight: 800; font-size: 16px; }}
.band .a .x:before {{ content: "→ "; color: {RED}; }}
.band.dark .a .x:before {{ color: {RED_L}; }}
.band .a .y {{ font-size: 13px; color: {MUTED}; margin-top: 4px; }}
.band.dark .a .y {{ color: #BDBDBD; }}
.mall {{ display: flex; gap: 16px; background: {SOFT}; border-radius: 16px; padding: 14px; margin-bottom: 12px; align-items: stretch; }}
.mall .nm {{ flex: 0 0 170px; background: {INK}; color: #fff; border-radius: 10px; padding: 14px 16px; display: flex; flex-direction: column; justify-content: center; }}
.mall .nm .t {{ font-weight: 800; font-size: 16px; }}
.mall .nm .r {{ font-size: 11.5px; color: #A8A8A8; margin-top: 8px; line-height: 1.6; }}
.mall .nm .r b {{ color: {RED_L}; }}
.mall .pairs {{ flex: 1; display: flex; flex-direction: column; justify-content: center; }}
.pair {{ display: flex; align-items: center; gap: 12px; padding: 7px 0; border-bottom: 1px dotted #D2D2D2; font-size: 13.5px; }}
.pair:last-child {{ border-bottom: 0; }}
.pair .p {{ flex: 1; color: {INK}; }}
.pair .ar {{ color: {RED}; font-weight: 900; }}
.pair .q {{ flex: 1.05; font-weight: 800; }}
.mall .ev {{ flex: 0 0 210px; background: #fff; border-radius: 10px; padding: 12px 14px; font-size: 12.5px; display: flex; flex-direction: column; justify-content: center; gap: 4px; }}
.mall .ev span {{ color: {MUTED}; margin-left: 6px; }}
.chip {{ display: inline-block; padding: 5px 12px; border-radius: 999px; border: 1px solid {LINE}; background: #fff; margin: 0 5px 6px 0; font-size: 12.5px; }}
.dark .chip {{ background: #333; border-color: #333; color: #fff; }}
.mini {{ font-size: 12px; color: {MUTED}; font-weight: 700; margin: 10px 0 6px; }}
.dark .mini {{ color: #A8A8A8; }}
.foc {{ font-weight: 800; font-size: 14.5px; line-height: 1.8; }}
.month {{ font-size: 30px; font-weight: 800; }}
.month.r {{ color: {RED}; }}
.tag {{ color: {MUTED}; font-weight: 700; margin-left: 10px; font-size: 14px; }}
.ul {{ margin: 6px 0 10px 0; padding-left: 18px; font-size: 14px; line-height: 1.7; }}
.done {{ border-radius: 14px; padding: 16px 20px; background: {SOFT}; font-weight: 800; font-size: 14px; line-height: 1.9; }}
.done.dark {{ background: {INK}; color: #fff; }}
.done .h {{ color: {RED}; font-size: 12px; }}
.done.dark .h {{ color: {RED_L}; }}
.note {{ font-size: 12.5px; color: {MUTED}; margin-top: 6px; }}
.callout {{ background: {INK}; color: #fff; border-radius: 14px; padding: 16px 22px; display: flex; gap: 22px; align-items: center; }}
.callout .l {{ color: {RED_L}; font-weight: 800; font-size: 16px; }}
.callout .t {{ font-weight: 800; font-size: 18px; }}
.callout .s {{ color: #BDBDBD; font-size: 13px; }}
div[data-testid="stMetric"] {{ background: {SOFT}; border-radius: 14px; padding: 14px 18px; }}
.stDownloadButton button {{ background: {RED}; border: 0; color: #fff; font-weight: 800; border-radius: 14px; padding: 12px 0; }}
.stDownloadButton button:hover {{ background: #A50D25; color: #fff; }}
.stLinkButton a {{ background: {INK}; border: 0; color: #fff !important; font-weight: 800; border-radius: 14px; padding: 10px 0; }}
.stLinkButton a:hover {{ background: #333; color: #fff !important; }}
div[data-testid="stMetricValue"] {{ font-weight: 800; font-size: 30px; letter-spacing: -.02em; }}
</style>
""", unsafe_allow_html=True)


def html(s: str):
    st.markdown("".join(line.strip() for line in s.splitlines()), unsafe_allow_html=True)


def header(kicker, title, sub=""):
    html(f'<div class="kicker">{kicker}</div><div class="h1">{title}</div>' + (f'<div class="sub">{sub}</div>' if sub else ""))


def eok(v):
    return f"{v / 1e8:,.1f}"


def plot_layout(fig, h=360, **kw):
    top = 56 if "title" in kw else 20
    if "title" in kw:
        kw["title"] = {**kw["title"], "x": 0, "xanchor": "left", "y": 0.98, "yanchor": "top", "pad": dict(l=6)}
    fig.update_layout(height=h, margin=dict(l=10, r=10, t=top, b=10), paper_bgcolor="rgba(0,0,0,0)", plot_bgcolor="rgba(0,0,0,0)",
                      font=dict(family=FONT, size=13, color=INK), showlegend=kw.pop("showlegend", False), **kw)
    fig.update_xaxes(showgrid=False, zeroline=False, linecolor=LINE)
    fig.update_yaxes(showgrid=True, gridcolor="#EAEAEA", zeroline=False)
    return fig


# ------------------------------------------------------------------ data
def find_file(keyword):
    if not DATA.exists():
        return None
    for p in sorted(DATA.glob("*.xls*")):
        if keyword in p.stem and not p.name.startswith("~$"):
            return p
    return None


@st.cache_data(show_spinner=False)
def load_targets(raw: bytes):
    """연간매출및 목표 시트: '병행'/'공식' 블록별 합계·몰별·월별 목표/실적."""
    wb = load_workbook(io.BytesIO(raw), data_only=True)
    ws = wb.worksheets[0]
    out = {}
    for r in range(1, ws.max_row + 1):
        label = ws.cell(r, 2).value
        if label not in ("병행", "공식") or label in out:
            continue
        # 월 헤더 위치: 같은 행에서 '1월'..'12월'
        months = {}
        for c in range(1, ws.max_column + 1):
            v = ws.cell(r, c).value
            if isinstance(v, str) and re.fullmatch(r"\d{1,2}월", v.strip()):
                months[int(v.strip()[:-1])] = c
        # 합계 행: 이후 첫 행 중 C열이 쇼핑몰/쇼핌몰 이고 E열이 숫자
        tot = None
        for rr in range(r + 1, r + 10):
            if str(ws.cell(rr, 3).value).strip() in ("쇼핑몰", "쇼핌몰") and isinstance(ws.cell(rr, 5).value, (int, float)):
                tot = rr
                break
        if tot is None:
            continue
        def row_dict(rr):
            d = {"몰": ws.cell(rr, 3).value, "목표": ws.cell(rr, 5).value or 0, "누적": ws.cell(rr, 6).value or 0}
            for m, c in months.items():
                t, a = ws.cell(rr, c).value, ws.cell(rr, c + 1).value
                d[f"{m}월_목표"] = t if isinstance(t, (int, float)) else 0
                d[f"{m}월_실제"] = a if isinstance(a, (int, float)) else None
            # 누적 = 1~9월 실제매출 합계 (파일 '누적매출' 칸은 10월 초 매출이 일부 포함돼 있어 사용하지 않음)
            d["누적"] = sum(d.get(f"{m}월_실제") or 0 for m in range(1, 10))
            return d
        total = row_dict(tot)
        malls = []
        for rr in range(tot + 1, ws.max_row + 1):
            name = ws.cell(rr, 3).value
            if name is None or ws.cell(rr, 2).value in ("병행", "공식"):
                break
            if isinstance(ws.cell(rr, 5).value, (int, float)):
                malls.append(row_dict(rr))
        out[label] = {"total": total, "malls": pd.DataFrame(malls)}
    return out


@st.cache_data(show_spinner=False)
def load_grades(raw: bytes):
    df = pd.read_excel(io.BytesIO(raw))
    df.columns = [str(c).strip() for c in df.columns]
    if "등급" not in df.columns:
        blank = [c for c in df.columns if c == "" or c.startswith("Unnamed")]
        if blank:
            df = df.rename(columns={blank[0]: "등급"})
    df["등급"] = df["등급"].astype(str).str.strip()
    if "재고" in df.columns:
        df = df[pd.to_numeric(df["재고"], errors="coerce").fillna(0) > 0].reset_index(drop=True)
    return df


@st.cache_data(show_spinner=False)
def to_xlsx(df: pd.DataFrame) -> bytes:
    buf = io.BytesIO()
    df.to_excel(buf, index=False, sheet_name="상품등급")
    return buf.getvalue()


with st.sidebar:
    st.image(str(BASE / "assets" / "logo_w.png"), width=170)
    st.markdown("<div style='height:6px'></div><div style='color:#8C8C8C;font-size:12px;letter-spacing:.12em'>2026 Q4 OPERATION PLAN</div>", unsafe_allow_html=True)
    st.markdown("<div style='height:18px'></div>", unsafe_allow_html=True)
    NAV = [("sec0", "표지"), ("sec1", "01  EC 목표매출"), ("sec2", "02  운영방향"), ("sec3", "03  채널 운영전략"),
           ("sec4", "04  브랜드 운영전략"), ("sec5", "05  체화재고 운영전략")]
    st.markdown('<nav class="toc">' + "".join(f'<a href="#{k}" target="_self">{v}</a>' for k, v in NAV) + "</nav>"
                + "<div class='toc-tip'>마우스 휠로 넘기면<br>다음 목차로 이동합니다</div>", unsafe_allow_html=True)
    # 휠 한 번 = 다음/이전 목차 (섹션이 화면보다 길면 그 안에서는 일반 스크롤)
    components.html("""<script>
const doc = window.parent.document;
function setup() {
  const main = doc.querySelector('[data-testid="stMain"]');
  if (!main) { setTimeout(setup, 300); return; }
  if (main.dataset.pager === "1") return;
  main.dataset.pager = "1";
  let busyUntil = 0, lastWheel = 0;
  main.addEventListener("wheel", (e) => {
    if (e.ctrlKey || Math.abs(e.deltaY) < Math.abs(e.deltaX)) return;
    const secs = [...doc.querySelectorAll('div[class*="st-key-sec"]')];
    if (!secs.length) return;
    const mt = main.getBoundingClientRect().top, vh = main.clientHeight;
    let cur = 0;
    secs.forEach((el, i) => { if (el.getBoundingClientRect().top - mt <= 40) cur = i; });
    const r = secs[cur].getBoundingClientRect();
    if (e.deltaY > 0 && r.bottom - mt > vh + 40) return;
    if (e.deltaY < 0 && r.top - mt < -40) return;
    e.preventDefault();
    const now = Date.now(), quiet = now - lastWheel > 220;
    lastWheel = now;
    if (now < busyUntil || (!quiet && now < busyUntil + 400)) return;
    const next = Math.min(secs.length - 1, Math.max(0, cur + (e.deltaY > 0 ? 1 : -1)));
    if (next === cur) return;
    busyUntil = now + 750;
    main.scrollTo({ top: main.scrollTop + secs[next].getBoundingClientRect().top - mt, behavior: "smooth" });
  }, { passive: false });
}
setup();
</script>""", height=0)

def read_bytes(upload, keyword):
    if upload is not None:
        return upload.getvalue(), upload.name
    p = find_file(keyword)
    return (p.read_bytes(), p.name) if p else (None, None)

t_raw, t_name = read_bytes(None, "목표")
g_raw, g_name = read_bytes(None, "등급")
TGT = load_targets(t_raw) if t_raw else None
GRD = load_grades(g_raw) if g_raw else None
b_raw, b_name = read_bytes(None, "브랜드")
BRD = pd.read_excel(io.BytesIO(b_raw)) if b_raw else None


# ------------------------------------------------------------------ static content
MAIN_MALLS = {  # 표시명: 목표매출 파일 몰명
    "신세계백화점몰": "신세계백화점몰", "신세계몰(SSG)": "신세계몰(신)", "롯데백화점": "롯데백화점온라인몰", "롯데홈쇼핑": "롯데홈쇼핑(신)",
    "현대홈쇼핑": "현대홈쇼핑(3)", "띵샵": "띵샵(신)", "삼성카드": "삼성카드쇼핑", "카카오톡선물하기": "카카오톡선물하기",
    "롯데쇼핑m몰": "롯데쇼핑(현대M몰)", "포이즌": "POIZON", "무신사 & 29cm": ["무신사", "29CM(공식)", "29CM"], "크림": "크림 주식회사", "W컨셉": "Wconcept",
    "에이블리": "에이블리", "지그재그": "카카오스타일 (지그재그)",
}
CH = {
    "신세계백화점몰": (["구찌 스카프·머플러", "AMI·비비안웨스트우드·캠퍼·메종키츠네·호카", "울프1834·파슬·아르마니"],
                ["멤버스 특가 활용 행사 운영", "쿠폰 조정·대형행사 오퍼로 가격 경쟁력 확보", "대형행사 메인 노출로 인지도 강화·재고소진"],
                [("쓱세일", "10/12~18"), ("쓱데이", "10/30~11/8")]),
    "신세계몰(SSG)": (["페라가모·몽블랑·프라다", "AMI·캠퍼·마르지엘라", "AI라이더스·보보쇼즈·미니로디니"],
                ["백화점 등록불가 브랜드 운영·차별화 구좌 확보", "모바일라이브·대형행사 쿠폰 협의 및 오퍼 활용", "유아동 대형행사 참여 및 단독행사 제안"],
                [("쓱세일", "10/12~18"), ("쓱데이", "10/30~11/8")]),
    "롯데백화점": (["버버리·프라다·구찌 등 가방", "에르노·막스마라·몽클레르 아우터", "아미,가니,비비안웨스트우드,메종마르지엘라"],
                ["고가 상품 백화점딜 운영", "아우터 및 부진 브랜드 행사 제안", "주력 판매브랜드로 매 행사시 물량소진"],
                [("L-PASS", "11/12~18"), ("광고 구좌", "매달 진행"), ("백화점딜", "상시")]),
    "롯데홈쇼핑": (["버버리·프라다·구찌", "톰브라운·스톤아일랜드·폴로 FW"], ["월별 대형 판촉 · 가방 메인 노출", "FW 적립 행사 및 연합전 제안"],
                [("광클절", "10/8~18"), ("대형 판촉", "월별")]),
    "현대홈쇼핑": (["몽클레르, 르메르·셀린느", "에르노·막스마라·버버리"], ["고가 가방/아우터 쇼핑라이브 예정", "입고 모델 주차별 라이브 선제안"],
                [("쇼핑라이브", "주차별")]),
    "띵샵": (["다미아니·프라다·톰브라운·구찌·론진·해밀턴", "판도라·폴로잡화·47브랜드"], ["무이자할부 활용 고단가 상품 집중 운영", "오픈런딜을 통한 재고 소진"],
                [("럭셔리 슈퍼위크", "매월"), ("연말 선물 기획", "12월")]),
    "삼성카드": (["라코스테, 호카·제이린드버그", "키츠네·AMI·비비안웨스트우드"], ["허용 범위 내 SS·FW 가격 대응", "부진 재고 및 과재고 SKU 상시 행사"],
                [("상시 행사", "")]),
    "카카오톡선물하기": (["카시오·톰브라운", "티켓투더문·아르마니·파슬"], ["몰 특성 맞는 카테고리·브랜드 행사가 운영", "쟁쟁한특가·포미위크 상시 참여로 재고소진"],
                [("카쇼페", "10/19~25"), ("쟁쟁한 특가", "월 1회")]),
    "롯데쇼핑m몰": (["라코스테, 구찌·CP컴퍼니·BARBOUR", "AMI·톰브라운·GANNI 등"], ["판매추이에 따른 쿠폰율 조정", "10~12월 비정기 기획전 추진"],
                [("비정기 기획전", "10~12월")]),
    "포이즌": (["폴로·BARBOUR·캠퍼", "클락스·노다·킨·어그"], ["주력 브랜드 상품 매핑 및 가격 점검", "잔여 부진 재고 소진"], []),
    "무신사 & 29cm": (["디젤, 티켓투더문", "파슬·아르마니"], ["디젤 익스클루시브 유지 및 무진장 추진", "티켓투더문 물량·노출구좌 확보"], [("무진장", "11/22~12/2"), ("이구데이(29cm)", "")]),
    "크림": (["RAB·노다·CEP·씨엘르", "CP컴퍼니·스톤아일랜드·보테가베네타"], ["스포츠 연합전 지속 제안", "인기 브랜드 가격·판매 반응 점검"], [("스포츠 연합전", "상시")]),
    "W컨셉": (["테클라·로이텀", "가니·리던·와일드동키·아페쎄·바버·파라부트", "파슬·비비안웨스트우드·아르마니"],
                ["리빙 행사 참여 및 연합전 노출", "프리미엄샵·더블유위크 매월 행사 참여", "월 2회 공식브랜드 기획전 진행"],
                [("프리미엄샵", "10/12~26"), ("더블유위크", "11/9~23"), ("윈터 페스타", "12/14~21")]),
    "에이블리": (["(4910) AMI·RAB·니들스·나나미카·호카·CP컴퍼니", "스와로브스키·판도라·비비안웨스트우드·아페쎄"],
                ["남성 타겟 차별화 브랜드·상품 판매", "월별 대형행사 수수료·쿠폰 협의로 가격 확보"],
                [("4910데이", "10/4~10"), ("글로벌위크", "10/9~18"), ("블랙프라이데이", "11/16~12/2"), ("연말결산", "12/23~1/1")]),
    "지그재그": (["판도라·카시오·비비안웨스트우드·아미", "파슬 잡화·비비안 시계·스부·어그·CK모자", "스와로브스키·아페쎄"],
                ["주력 브랜드 수수료 조정으로 가격 메리트 확보", "상품군 확대로 신규 매출 확보", "부진 브랜드 가격 경쟁으로 적극 운영"],
                [("직잭팟", "9/28~10/12"), ("겨울 블프", "11/16~12/2")]),
}
GROUPS = {
    "종합몰": (["신세계백화점몰", "신세계몰(SSG)", "롯데백화점", "롯데홈쇼핑", "현대홈쇼핑"], ["쓱세일·쓱데이 등 대형 행사 연계", "FW 의류·잡화 단품 행사"], "🏬"),
    "폐쇄몰": (["띵샵", "삼성카드", "카카오톡선물하기", "롯데쇼핑m몰", "포이즌"], ["카쇼페·슈퍼위크 등 몰 특화 행사", "대물량·고단가 상품 특가", "공식브랜드 패밀리세일 운영"], "🔒"),
    "패션 플랫폼": (["무신사 & 29cm", "크림", "W컨셉", "에이블리", "지그재그"], ["플랫폼별 대형행사 참여", "수수료·쿠폰 협의로 가격 경쟁력"], "👕"),
}
GROUP_NOTES = {"폐쇄몰": "라코스테 가격·쿠폰은 브랜드 정책 확인 후 적용", "패션 플랫폼": "디젤의 타 채널 배정은 익스클루시브 조건 확인 필요"}


def parse_period(s):
    """'10/12~18', '10/30~11/8', '12/23~1/1' -> (start, end) or None"""
    m = re.fullmatch(r"(\d{1,2})/(\d{1,2})~(?:(\d{1,2})/)?(\d{1,2})", s.replace(" ", ""))
    if not m:
        return None
    m1, d1, m2, d2 = int(m.group(1)), int(m.group(2)), m.group(3), int(m.group(4))
    m2 = int(m2) if m2 else m1
    y1 = 2026
    y2 = 2027 if m2 < m1 else 2026
    return date(y1, m1, d1), date(y2, m2, d2)


def mall_rates(display):
    if TGT is None:
        return ""
    key = MAIN_MALLS.get(display)
    keys = key if isinstance(key, list) else [key]
    parts = []
    for lab in ("병행", "공식"):
        df = TGT[lab]["malls"]
        rows = df[df["몰"].isin(keys)]
        tgt, acc = rows["목표"].sum(), rows["누적"].sum()
        if tgt:
            parts.append(f"{lab} <b>{acc / tgt * 100:.0f}%</b>")
    return " · ".join(parts)


BRAND_SHORT = {  # 브랜드 엑셀 '4분기 브랜드 운영 전략 코멘트' 요약 (없는 브랜드는 원문 표시)
    "AMI": "전 채널 주력 운영 · 단품 구좌 확보", "구찌": "스카프 중심 폐쇄몰 적극 운영",
    "비비안웨스트우드": "지그재그·에이블리 메인 · 시계 자체 블프", "페라가모": "신세계몰 메인 브랜드 운영",
    "헬렌카민스키": "4분기 역시즌 특가 운영", "라코스테": "시즌오프 전 M몰·이지웰·삼성카드 가격 대응",
    "GANNI": "주요 채널 집중 · SS 과재고 역시즌 특가", "메종마르지엘라": "전 채널 주력 · 부진 잡화 제안",
    "아페쎄": "지그재그·에이블리 가격 경쟁 운영", "아르마니": "쿠팡 배너·와우데이 · 카톡 균일가 행사",
}


# ------------------------------------------------------------------ pages
with st.container(key="sec0"):
    html('<div id="sec0" class="anchor"></div>')
    html("""
<div class="hero">
<div class="q4">Q4</div>
<div class="k">2026 Q4 OPERATION PLAN</div>
<div class="t">2026년 4분기<br>운영계획 및 매출 확대 방안</div>
<div class="s">e-커머스 운영 · 채널별 판매 확대 및 재고 소진 계획</div>
</div>""")
    st.markdown("<div style='height:16px'></div>", unsafe_allow_html=True)
    html('<div class="sec">목차</div>')
    for n, t, d in [("01", "EC 목표매출", "EC 전체 목표매출 · 공식 목표매출 변경 반영"),
                    ("02", "운영방향", "등급 기반 상품 회전 · 폐쇄몰 특가 · 업무 자동화 · 신규 채널 확대"),
                    ("03", "채널 운영전략", "행사일정 · 주력브랜드 정리 — 종합몰 / 폐쇄몰 / 패션 플랫폼"),
                    ("04", "브랜드 운영전략", "재고액순"), ("05", "체화재고 운영전략", "입고일순")]:
        html(f'<div class="agenda"><div class="n">{n}</div><div class="t">{t}</div><div class="d">{d}</div></div>')

with st.container(key="sec1"):
    html('<div id="sec1" class="anchor"></div>')
    header("01  EC 목표매출", "EC 전체 목표매출 및 진행 현황", "2026년 9월 누적 기준 · EC 전체 · 병행과 공식은 별도 집계하며 합산하지 않음")
    if not TGT:
        st.warning("목표매출 엑셀이 없습니다. data 폴더에 넣어 주세요.")
        st.stop()
    cols = st.columns(2, gap="medium")
    for i, lab in enumerate(("병행", "공식")):
        t = TGT[lab]["total"]
        rate = t["누적"] / t["목표"] * 100
        q4 = sum(t.get(f"{m}월_목표", 0) for m in (10, 11, 12))
        dark = i == 0
        with cols[i]:
            html(f"""
<div class="card {'dark' if dark else ''}">
<span class="pill {'red' if dark else ''}">{lab}</span><span style="margin-left:12px;color:{'#A8A8A8' if dark else MUTED};font-size:14px">현재 달성률</span>
<div class="big">{rate:.1f}<small>%</small></div>
<div class="track"><div class="fill" style="width:{min(rate, 100):.1f}%"></div></div>
<div class="stats">
<div><div class="l">2026년 총 목표</div><div class="v">{eok(t['목표'])}<span>억 원</span></div></div>
<div><div class="l">누적 매출</div><div class="v">{eok(t['누적'])}<span>억 원</span></div></div>
<div><div class="l">4분기 목표</div><div class="v r">{eok(q4)}<span>억 원</span></div></div>
</div></div>""")

    html('<div class="sec">월별 목표 대비 실적</div>')
    sel = st.radio("구분", ["병행", "공식"], horizontal=True, label_visibility="collapsed", key="m_sel")
    t = TGT[sel]["total"]
    ms = list(range(1, 13))
    tgt = [t.get(f"{m}월_목표", 0) / 1e8 for m in ms]
    act = [(t.get(f"{m}월_실제") or 0) / 1e8 if m <= 9 else None for m in ms]
    fig = go.Figure()
    fig.add_bar(x=[f"{m}월" for m in ms], y=[a if a is not None else 0 for a in act], name="실제 매출",
                marker_color=[INK if m <= 9 else "rgba(0,0,0,0)" for m in ms],
                text=[f"{a:.1f}" if a else "" for a in act], textposition="outside")
    fig.add_bar(x=[f"{m}월" for m in ms], y=[tgt[i] if ms[i] >= 10 else 0 for i in range(12)], name="4분기 목표",
                marker_color=RED, text=[f"{tgt[i]:.1f}" if ms[i] >= 10 else "" for i in range(12)], textposition="outside")
    fig.add_scatter(x=[f"{m}월" for m in ms], y=tgt, name="월 목표", mode="lines+markers", line=dict(color=MID, width=2, dash="dot"))
    plot_layout(fig, 300, barmode="overlay", showlegend=True, legend=dict(orientation="h", y=1.1, x=0))
    st.plotly_chart(fig, use_container_width=True, config={"displayModeBar": False})

    html(f'<div class="note">출처: {t_name} · 누적 매출 = 1~9월 실제매출 합계</div>')

with st.container(key="sec2"):
    html('<div id="sec2" class="anchor"></div>')
    header("02  운영방향", "운영방향", "등급 기반 상품 회전 · 폐쇄몰 특가 · 업무 자동화 · 신규 채널 확대")
    items = [("🏅", "등급 기반 상품 회전 관리", "자체등급 산정 후 상품 회전에 반영<br>B·C 행사·노출 집중 / D·E·F 가격 조정 강화"),
             ("🎁", "폐쇄몰 공식브랜드 특가운영", "자체 패밀리세일로<br>공식브랜드 특가 운영"),
             ("🤖", "업무 자동화", "반복 업무 자동화로<br>AMD 인력을 MD 업무에 투입"),
             ("🏪", "신규 채널 확대", "Npay 복지몰(비즈마켓) · 해외몰 입점 확대<br>등급별 재고의 추가 판로 확보")]
    cs = st.columns(4, gap="small")
    for n, (c, (ic, t, d)) in enumerate(zip(cs, items), start=1):
        with c:
            html(f'<div class="card" style="height:160px"><span class="num" style="font-size:20px">0{n}</span><div class="dot" style="width:40px;height:40px;font-size:19px">{ic}</div>'
                 f'<div class="dir" style="display:block;margin-top:10px"><div class="t" style="font-size:15px">{t}</div><div class="d" style="font-size:12.5px;line-height:1.55">{d}</div></div></div>')
    st.markdown("<div style='height:14px'></div>", unsafe_allow_html=True)
    html('<div style="display:flex;align-items:baseline;gap:12px;margin:2px 0 4px"><span class="kicker" style="margin:0">02-1  상품 등급 · 회전</span><span style="font-size:20px;font-weight:800">상품 등급 기반 회전 전략</span></div>')
    if GRD is None:
        st.warning("상품등급 엑셀이 없습니다. data 폴더에 넣어 주세요.")
        st.stop()
    order = ["S", "A", "B", "C", "D", "E", "F"]
    g = GRD[GRD["등급"].isin(order)]
    agg = g.groupby("등급").agg(lines=("라인명", "count"), sales=("매출", "sum"), margin=("이익율(%)", "median")).reindex(order)
    total = int(agg["lines"].sum())
    agg["share"] = agg["sales"] / agg["sales"].sum() * 100
    html(f'<div class="sub" style="margin-bottom:8px;font-size:13px">상품등급 10/8 기준 · 재고 보유 {total:,}개 라인 (재고 0 제외) · {g["브랜드"].nunique()}개 브랜드 · S~F 7단계 등급을 상품 회전 전략에 반영</div>')
    html('<div class="steps">' + '<span class="chev">›</span>'.join(
        f'<div class="step {"on" if i == 0 else ""}"><b>0{i + 1}</b>{s}</div>' for i, s in
        enumerate(["상품 등급 산정 (S~F)", "등급별 회전 전략 수립", "행사·노출 / 가격 조정 반영", "판매 결과로 등급 갱신"])) + '</div>')

    color = {"S": INK, "A": INK, "B": MID, "C": MID, "D": RED, "E": RED, "F": RED}
    left, right = st.columns([4, 8], gap="medium")
    with left:
        fig = go.Figure(go.Bar(x=order, y=agg["lines"], marker_color=[color[k] for k in order],
                               text=[f"{int(v):,}" for v in agg["lines"]], textposition="outside",
                               customdata=agg[["share", "margin"]].values,
                               hovertemplate="%{x}등급<br>라인 %{y:,}개<br>매출 비중 %{customdata[0]:.1f}%<br>이익율 중앙값 %{customdata[1]:.1f}%<extra></extra>"))
        plot_layout(fig, 255, title=dict(text="등급별 라인 수", font=dict(size=15)))
        fig.update_yaxes(visible=False)
        fig.update_xaxes(tickfont=dict(size=15))
        st.plotly_chart(fig, use_container_width=True, config={"displayModeBar": False})
        st.link_button("🔎  상품 수익율 검색기 바로가기", "https://search-data-cga2d6ch3cz3rufnlhmsih.streamlit.app/", use_container_width=True)
    with right:
        def band(keys):
            b = agg.loc[keys]
            return int(b["lines"].sum()), b["lines"].sum() / total * 100, b["share"].sum(), b["margin"].min(), b["margin"].max()
        for lab, nm, keys, bc, dark, act, sub in [
            ("S·A", "핵심 상품", ["S", "A"], INK, False, "재고 우선 확보 · 정상가 판매 유지", "주력 채널 메인 구좌 우선 배정"),
            ("B·C", "중위 상품", ["B", "C"], MID, False, "행사 참여 · 노출 구좌 집중", "대형 행사·기획전 노출로 판매 속도 제고"),
            ("D·E·F", "비인기 상품", ["D", "E", "F"], RED, True, "가격 조정 강화로 빠른 회전", "체화재고 → 자체 클리어런스 운영<br>시즌아웃 재고 → 기존 시즌행사가 유지<br><span style='opacity:.75;font-size:12px'>기존엔 다음 시즌 대비 가격 원복(상승) → 체화재고 누적으로 시즌오프 행사가 유지</span>")]:
            n, lp, sp, mn, mx = band(keys)
            mt = f"{mx:.0f}% 이하" if lab == "D·E·F" else f"{mn:.0f}~{mx:.0f}%"
            html(f"""<div class="band {'dark' if dark else ''}"><div class="b" style="background:{bc}">{lab}</div>
<div class="m"><div class="n">{nm}</div><div class="s">{n:,}개 라인 ({lp:.0f}%)<br>매출 {sp:.1f}% · 이익율 {mt}</div></div>
<div class="a"><div class="x">{act}</div><div class="y">{sub}</div></div></div>""")
        st.download_button("⬇  상품등급 전체 목록 내려받기 (엑셀 · 재고 보유)", to_xlsx(GRD), file_name="상품등급_재고보유_1008.xlsx",
                           mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet", type="primary", use_container_width=True)


with st.container(key="sec3"):
    html('<div id="sec3" class="anchor"></div>')
    header("03  채널 운영전략", "채널 운영전략", "주요 16개 몰 — 종합몰 · 폐쇄몰 · 패션 플랫폼별 행사일정 · 주력브랜드 정리")
    cs = st.columns(3, gap="medium")
    for i, (gname, (malls, focus, ic)) in enumerate(GROUPS.items()):
        dark = True
        with cs[i]:
            html(f"""<div class="card {'dark' if dark else ''}" style="height:284px;display:flex;flex-direction:column;box-sizing:border-box"><div style="display:flex;gap:12px;align-items:center"><div class="dot" style="flex:0 0 46px;height:46px">{ic}</div>
<div style="font-size:22px;font-weight:800">{gname}</div></div><div class="mini">주요 채널</div><div>{''.join(f'<span class="chip">{c}</span>' for m in malls for c in m.split(' & '))}</div>
<div style="margin-top:auto"><div class="mini">4분기 집중 방향</div><div class="foc" style="min-height:66px;font-size:13.5px;line-height:1.65">{'<br>'.join('▪ ' + f for f in focus)}</div></div></div>""")
    st.markdown("<div style='height:14px'></div>", unsafe_allow_html=True)

    tabs = st.tabs(["📅 행사 캘린더", "🏬 종합몰", "🔒 폐쇄몰", "👕 패션 플랫폼", "🗓 10~12월 운영"])
    with tabs[0]:
        rows = []
        for gname, (malls, _, _) in GROUPS.items():
            for m in malls:
                for ev, d in CH[m][2]:
                    p = parse_period(d) if d else None
                    if p:
                        rows.append(dict(몰=m, 행사=ev, 시작=p[0], 종료=p[1], 그룹=gname))
        ev = pd.DataFrame(rows).sort_values(["몰", "시작"])
        ev["y"] = ev["몰"]  # 겹치는 행사도 몰별 한 줄에 표시
        order_m = [y for g_ in GROUPS.values() for m in g_[0] for y in sorted(set(ev.loc[ev["몰"] == m, "y"]), key=len)]
        gcol = {"종합몰": INK, "폐쇄몰": MUTED, "패션 플랫폼": RED}
        OVER = "#F2A900"  # 실제 기간이 다른 행사와 겹치는 경우 테두리
        PX_DAY = 8.5      # 대략 1일 = 8.5px → 글자가 다 들어가도록 막대 최소 길이 계산
        def label_days(t):
            return (sum(11.5 if ord(c) > 0x3000 else 7 for c in t) + 22) / PX_DAY
        D1 = pd.Timedelta(days=1)
        fig = go.Figure()
        for m_, grp in ev.groupby("몰", sort=False):
            grp = grp.sort_values("시작").reset_index(drop=True)
            real = [(pd.Timestamp(r["시작"]), pd.Timestamp(r["종료"]) + D1) for _, r in grp.iterrows()]
            prev_end = None
            for i, r in grp.iterrows():
                s0, e0 = real[i]
                overlaps = any(k != i and real[k][0] < e0 and s0 < real[k][1] for k in range(len(real)))
                width = max(e0 - s0, pd.Timedelta(days=label_days(r["행사"])))
                bs = s0 + (e0 - s0) / 2 - width / 2
                if prev_end is not None and bs < prev_end + pd.Timedelta(hours=14):
                    bs = prev_end + pd.Timedelta(hours=14)
                be = bs + width
                prev_end = be
                fig.add_trace(go.Bar(base=[bs], x=[(be - bs).total_seconds() * 1000], y=[m_], orientation="h",
                                     marker=dict(color=gcol[r["그룹"]], line=dict(color=OVER if overlaps else gcol[r["그룹"]], width=3 if overlaps else 0)),
                                     hovertemplate=f"{m_} · {r['행사']}<br>{r['시작']:%m/%d} ~ {r['종료']:%m/%d}<extra></extra>"))
                fig.add_annotation(x=bs + width / 2, y=m_, text=r["행사"], showarrow=False, font=dict(color="#fff", size=12))
        plot_layout(fig, 40 + 29 * len(order_m), barmode="overlay", bargap=0.28)
        fig.update_xaxes(type="date", range=["2026-09-22", "2027-01-05"], dtick="M1", tickformat="%m월", showgrid=True, gridcolor="#EAEAEA", side="top")
        fig.update_yaxes(categoryorder="array", categoryarray=order_m[::-1], showgrid=False)
        st.plotly_chart(fig, use_container_width=True, config={"displayModeBar": False})
        html('<div class="note">■ 종합몰 · <span style="color:#6E6E6E">■</span> 폐쇄몰 · <span style="color:#C8102E">■</span> 패션 플랫폼 · <span style="color:#F2A900">□</span> 노란 테두리 = 기간이 겹치는 행사 · 막대 길이는 글자에 맞춘 대략적 기간 (정확한 날짜는 마우스 오버) · 날짜가 정해진 행사만 표시 (상시·월별 행사는 채널 탭 참고)</div>')

    def mall_cards(gname):
        for m in GROUPS[gname][0]:
            bl, al, evs = CH[m]
            pairs = "".join(f'<div class="pair"><div class="p">{b}</div><div class="ar">➜</div><div class="q">{a}</div></div>' for b, a in zip(bl, al))
            evh = "".join(f'<div><b>{e}</b><span>{d}</span></div>' for e, d in evs) or '<div style="color:#9A9A9A">—</div>'
            rates = mall_rates(m)
            html(f"""<div class="mall"><div class="nm"><div class="t">{m}</div>{f'<div class="r">9월 누적 달성률<br>{rates}</div>' if rates else ''}</div>
<div class="pairs">{pairs}</div><div class="ev">{evh}</div></div>""")
        if gname in GROUP_NOTES:
            html(f'<div class="note">ⓘ {GROUP_NOTES[gname]}</div>')

    for tab, gname in zip(tabs[1:4], GROUPS):
        with tab:
            html('<div style="display:flex;gap:16px;font-size:12px;color:#6E6E6E;font-weight:700;margin:4px 0 8px"><div style="flex:0 0 170px;padding-left:14px">채널</div><div style="flex:1">주력 브랜드 ➜ 판매 확대 방안</div><div style="flex:0 0 210px">행사 일정</div></div>')
            mall_cards(gname)
    with tabs[4]:
        cs = st.columns(3, gap="medium")
        data = [("10월", "준비 · 선정", ["FW 주력 SKU·장기 재고 선정", "스카프 집중 판매 및 단품 행사 제안", "롯데홈쇼핑 광클절"], ["채널별 상품·판매가·물량 확정", "11월 행사 조건 및 구좌 협의"]),
                ("11월", "집중 판매", ["모든 쇼핑몰 블랙프라이데이 및 무진장 참여 추진", "주력 상품 물량 집중", "신세계V 슥데이 동시진행"], ["행사별 추가 매출·수익성 확인", "판매 반응에 따라 후속 행사 확대"]),
                ("12월", "잔여 소진 · 마감", ["겨울 잔여 재고·연말 선물 행사", "미달 채널 추가 프로모션"], ["잔여 재고별 후속 계획 실행", "연말 실적 및 미소진 상품 정리"])]
        for i, (c, (mo, tag, ops, done)) in enumerate(zip(cs, data)):
            with c:
                html(f"""<div><span class="month {'r' if i == 1 else ''}">{mo}</span><span class="tag">{tag}</span></div>
<div class="mini">주요 운영</div><ul class="ul">{''.join(f'<li>{o}</li>' for o in ops)}</ul>
<div class="done {'dark' if i == 1 else ''}"><div class="h">완료 기준</div>{'<br>'.join('✓ ' + d for d in done)}</div>""")

with st.container(key="sec4"):
    html('<div id="sec4" class="anchor"></div>')
    hl, hr = st.columns([7.5, 2.5], vertical_alignment="center")
    with hl:
        header("04  브랜드 운영전략", "브랜드 운영전략 (재고액순)", "재고액 TOP 10 · 재고액과 당월 매출을 함께 검토해 판매 기회 확보")
    with hr:
        if b_raw:
            st.download_button("⬇  전 브랜드 실행방향 (엑셀)", b_raw, file_name=b_name or "브랜드재고.xlsx",
                               mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet", type="primary", use_container_width=True)
    if BRD is None:
        st.warning("브랜드 재고 엑셀이 없습니다. data 폴더에 파일명에 '브랜드'가 들어간 엑셀을 넣어 주세요.")
        st.stop()
    top = BRD.sort_values("재고액", ascending=False).head(10)
    br = pd.DataFrame({"브랜드": top["브랜드"], "재고액(억)": top["재고액"] / 1e8, "당월 매출(만원)": (top["당월 매출"] / 1e4).round().astype(int),
                       "우선 실행 방향": [BRAND_SHORT.get(b, c) for b, c in zip(top["브랜드"], top["4분기 브랜드 운영 전략 코멘트"])]})
    l, r = st.columns([4, 8], gap="medium")
    with l:
        fig = go.Figure(go.Bar(x=br["재고액(억)"][::-1], y=br["브랜드"][::-1], orientation="h", marker_color=RED,
                               text=[f"{v:.2f}" for v in br["재고액(억)"][::-1]], textposition="outside", cliponaxis=False))
        fig.update_xaxes(range=[0, br["재고액(억)"].max() * 1.15])
        plot_layout(fig, 400, title=dict(text="브랜드별 재고액 (억 원)", font=dict(size=14)))
        fig.update_xaxes(visible=False)
        st.plotly_chart(fig, use_container_width=True, config={"displayModeBar": False})
    with r:
        st.dataframe(br, hide_index=True, use_container_width=True, height=388,
                     column_config={"브랜드": st.column_config.TextColumn(width=130),
                                    "재고액(억)": st.column_config.NumberColumn(format="%.2f", width=85),
                                    "당월 매출(만원)": st.column_config.NumberColumn(format="%,d", width=115),
                                    "우선 실행 방향": st.column_config.TextColumn(width=400)})
        if GRD is not None:
            exp = st.expander("브랜드별 등급 분포 (상품등급 기준)", expanded=False)
            sub = GRD[GRD["브랜드"].isin(br["브랜드"])]
            if len(sub):
                pv = sub.pivot_table(index="브랜드", columns="등급", values="라인명", aggfunc="count", fill_value=0)
                pv = pv.reindex(columns=[c for c in ["S", "A", "B", "C", "D", "E", "F"] if c in pv.columns])
                exp.dataframe(pv.reindex([b for b in br["브랜드"] if b in pv.index]), use_container_width=True)
            else:
                st.caption("상품등급 파일에서 해당 브랜드명을 찾지 못했습니다.")

with st.container(key="sec5"):
    html('<div id="sec5" class="anchor"></div>')
    header("05  체화재고 운영전략", "체화재고 운영전략 (입고일순)", "FW 장기 재고와 시즌이 지난 상품을 별도 운영")
    cards = [("❄️", "FW 장기 재고", "톰브라운, 제이린드버그<br>텐씨, ADD, 에르노, 피레넥스 등", ["10월 상품 선정", "11월 집중 행사", "12월 잔여 사이즈별 후속 제안"]),
             ("👜", "잡화 과재고", "지방시 스카프<br>마르니·마르지엘라·발렌티노 잡화", ["단품 행사 구좌 확보", "지방시 스카프 10~11월 집중", "지속적 납품 제안 (쿠팡·아이몰 등)"]),
             ("☀️", "역시즌 상품", "헬렌카민스키<br>GANNI의 SS 과재고", ["겨울 상품과 분리해 특가 제안", "채널별 시즌오프 행사 운영 가능 여부 확인"]),
             ("🎚️", "판매 조건 관리", "스톤아일랜드 / 호카", ["스톤아일랜드 마진 유지", "호카 특정 사이즈 구매 제한 검토"])]
    cs = st.columns(4, gap="medium")
    for i, (c, (ic, t, tg, plan)) in enumerate(zip(cs, cards)):
        with c:
            html(f"""<div class="card" style="height:430px;display:flex;flex-direction:column"><div class="dot" style="width:52px;background:{RED if i == 0 else INK}">{ic}</div>
<div style="font-size:19px;font-weight:800;margin:16px 0 12px">{t}</div><div class="mini">대상 브랜드 / 상품</div><div style="font-size:14px;line-height:1.7">{tg}</div>
<div style="background:#fff;border-radius:12px;padding:14px 16px;margin-top:auto;height:165px;box-sizing:border-box"><div style="color:{RED};font-weight:800;font-size:12px;margin-bottom:6px">실행 계획</div>
<div style="font-weight:800;font-size:13.5px;line-height:1.9">{'<br>'.join('▪ ' + p for p in plan)}</div></div></div>""")
